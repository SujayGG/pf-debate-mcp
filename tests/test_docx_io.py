from pf_debate_mcp.cards import runs_from_markup
from pf_debate_mcp.docx_io import export, parse

CARD = {
    "tag": "Moratorium kills US AI lead",
    "cite_short": "Smith 25",
    "cite_rest": "[Jane Smith; fellow at CSIS; \"Compute\" Foreign Affairs; 2025]",
    "runs": [("Compute is ", False, False), ("the binding constraint", True, True),
             (" on frontier AI, and ", True, False), ("China is catching up", True, True), (".", False, False)],
}


def test_round_trip(tmp_path):
    path = export("Aff Case", [{"pocket": "Aff"}, {"block": "C1: AI"}, {"card": CARD}, {"text": "Analytic."}],
                  tmp_path / "case.docx")
    cards = parse(path.read_bytes())
    assert len(cards) == 1
    c = cards[0]
    assert (c["pocket"], c["block"], c["tag"], c["cite_short"]) == ("Aff", "C1: AI", CARD["tag"], "Smith 25")
    assert c["cite_rest"] == CARD["cite_rest"]
    assert c["runs"] == CARD["runs"]


def test_blockfile_layout_round_trip(tmp_path):
    items = [{"text": "Resolved: something."}, {"pocket": "PRO / A2 CON"}, {"hat": "A2 Cyberattacks"},
             {"card": {**CARD, "tag": "1. [NU] Breaches already triple"}}, {"text": "- Risk is the baseline"},
             {"tag": "2. [DL] Not a backdoor (analytic)"}, {"text": "- Court order first"}]
    path = export("PESH MASTER BLOCKFILE", items, tmp_path / "bf.docx")
    from docx import Document
    doc = Document(str(path))
    assert doc.paragraphs[0].style.name == "Title" and doc.paragraphs[0].text == "PESH MASTER BLOCKFILE"
    assert [p.text for p in doc.paragraphs if p.style.name == "List Bullet"] == ["Risk is the baseline",
                                                                                 "Court order first"]
    read = [r for p in doc.paragraphs for r in p.runs if r.text == "the binding constraint"][0]
    assert read.bold and read.underline and read.font.size.pt == 12
    cards = parse(path.read_bytes())
    assert [(c["pocket"], c["hat"], c["tag"]) for c in cards] == [("PRO / A2 CON", "A2 Cyberattacks",
                                                                   "1. [NU] Breaches already triple")]
    assert cards[0]["runs"] == CARD["runs"]  # the warrant bullet is not folded into the card


def test_dataset_markup():
    m = "<h4>Tag</h4><p><strong>X 12</strong> cite</p><p>Lastly, <u><mark>weak <strong>econ</strong></mark>. Other</u> text</p>"
    assert runs_from_markup(m) == [("Lastly, ", False, False), ("weak econ", True, True),
                                   (". Other", True, False), (" text", False, False)]
