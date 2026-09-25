"""Verbatim-compatible .docx export and parsing.

Verbatim (the standard debate Word template) maps Pocket/Hat/Block/Tag onto Word's
Heading 1-4, so plain Heading styles open correctly in Verbatim and in plain Word.
"""

import html
import io
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.shared import Pt, RGBColor

from .cards import Run, normalize_runs

_HEADINGS = {"pocket": (1, 26), "hat": (2, 22), "block": (3, 16), "tag": (4, 13)}


def _style_doc(doc) -> None:
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    for level, size in _HEADINGS.values():
        st = doc.styles[f"Heading {level}"]
        st.font.size, st.font.bold, st.font.name = Pt(size), True, "Calibri"
        st.font.color.rgb = RGBColor(0, 0, 0)
        if level <= 2:
            st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # Speech prose between cards gets its own style so parse() never folds it into a card body.
    doc.styles.add_style("Analytic", WD_STYLE_TYPE.PARAGRAPH).base_style = doc.styles["Normal"]


def export(title: str, items: list[dict], path: Path) -> Path:
    """items: {"pocket"|"hat"|"block"|"tag": str} | {"text": str} | {"card": card dict}.
    A text starting with "- " is a bullet (an analytic warrant under a response)."""
    doc = Document()
    _style_doc(doc)
    doc.core_properties.title = title
    doc.add_paragraph(title, style="Title")
    for item in items:
        if "card" in item:
            c = item["card"]
            doc.add_heading(c["tag"], 4)
            p = doc.add_paragraph()
            short = p.add_run(c["cite_short"])
            short.bold, short.font.size = True, Pt(13)
            p.add_run(" " + (c.get("cite_rest") or "")).font.size = Pt(8)
            p = doc.add_paragraph()
            for text, u, h in c["runs"]:
                if text == "\n":
                    p = doc.add_paragraph()
                    continue
                r = p.add_run(text)
                if h:  # read aloud: bold, underlined, highlighted, full size
                    r.bold = r.underline = True
                    r.font.size = Pt(12)
                    r.font.highlight_color = WD_COLOR_INDEX.TURQUOISE
                elif u:  # context
                    r.underline = True
                    r.font.size = Pt(10)
                else:
                    r.font.size = Pt(8)  # unread text is shrunk, the Verbatim convention
        elif "text" in item:
            t = item["text"]
            if t.startswith("- "):
                doc.add_paragraph(t[2:], style="List Bullet")
            else:
                doc.add_paragraph(t, style="Analytic")
        else:
            (kind, text), = item.items()
            doc.add_heading(text, _HEADINGS[kind][0])
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


def export_html(title: str, items: list[dict], path: Path) -> Path:
    """Same structure as export() as HTML, which Google Docs keeps intact on paste or Drive upload
    (headings, bold cite, underline, highlight, small unread text)."""
    esc = html.escape
    out = [f"<!doctype html><meta charset='utf-8'><title>{esc(title)}</title>"
           "<body style='font-family:Calibri,Arial,sans-serif;font-size:11pt'>"
           f"<p style='font-size:26pt'>{esc(title)}</p>"]
    for item in items:
        if "card" in item:
            c = item["card"]
            out.append(f"<h4>{esc(c['tag'])}</h4><p><b style='font-size:13pt'>{esc(c['cite_short'])}</b> "
                       f"<span style='font-size:8pt'>{esc(c.get('cite_rest') or '')}</span></p><p>")
            for text, u, h in c["runs"]:
                t = esc(text).replace("\n", "</p><p>")
                if h:
                    t = f"<b><u style='background-color:#00ffff;font-size:12pt'>{t}</u></b>"
                elif u:
                    t = f"<u style='font-size:10pt'>{t}</u>"
                else:
                    t = f"<span style='font-size:8pt'>{t}</span>"
                out.append(t)
            out.append("</p>")
        elif "text" in item:
            t = item["text"]
            out.append(f"<ul><li>{esc(t[2:])}</li></ul>" if t.startswith("- ") else f"<p>{esc(t)}</p>")
        else:
            (kind, text), = item.items()
            out.append(f"<h{_HEADINGS[kind][0]}>{esc(text)}</h{_HEADINGS[kind][0]}>")
    out.append("</body>")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("".join(out), encoding="utf-8")
    return path


def _level(paragraph) -> str | None:
    name = (paragraph.style.name or "").lower()
    for key, (level, _) in _HEADINGS.items():
        if name == f"heading {level}" or key in name:
            return key
    return None


MAX_DOCX, MAX_UNZIPPED, MAX_ENTRIES, MAX_CARDS = 10_000_000, 100_000_000, 2_000, 500


def parse(data: bytes) -> list[dict]:
    """Split a Verbatim-style speech doc into cards with their pocket/hat/block context.
    Untrusted input: size limits are checked on the zip directory before anything is decompressed."""
    if len(data) > MAX_DOCX:
        raise ValueError("document is over 10 MB")
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        entries = z.infolist()
        if len(entries) > MAX_ENTRIES or sum(e.file_size for e in entries) > MAX_UNZIPPED:
            raise ValueError("document expands to an unreasonable size (possible zip bomb)")
    doc = Document(io.BytesIO(data))
    ctx = {"pocket": None, "hat": None, "block": None}
    cards, cur = [], None
    for p in doc.paragraphs:
        level = _level(p)
        if level and level != "tag":
            ctx[level] = p.text.strip()
            if level == "pocket":
                ctx["hat"] = ctx["block"] = None
            elif level == "hat":
                ctx["block"] = None
            cur = None
            continue
        if (p.style.name or "").lower() in ("analytic", "list bullet", "title"):
            cur = None
            continue
        if level == "tag":
            cur = {**ctx, "tag": p.text.strip(), "cite_short": None, "cite_rest": "", "runs": []}
            cards.append(cur)
            continue
        if cur is None or not p.text.strip():
            continue
        if cur["cite_short"] is None:
            bold = "".join(r.text for r in p.runs if r.bold).strip()
            cur["cite_short"] = bold or p.text.strip()[:40]
            cur["cite_rest"] = p.text.strip()[len(bold):].strip() if bold else p.text.strip()
            continue
        if cur["runs"]:
            cur["runs"].append(("\n", False, False))
        for r in p.runs:
            sname = (r.style.name or "").lower() if r.style is not None else ""
            u = bool(r.underline) or "underline" in sname or "emphasis" in sname
            h = r.font.highlight_color is not None
            cur["runs"].append((r.text, u, h))
    out = []
    for c in cards:
        if c["runs"]:
            c["runs"] = normalize_runs(c["runs"])
            out.append(c)
    return out[:MAX_CARDS]
